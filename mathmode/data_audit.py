"""Deterministically inspect original JSON/CSV/document inputs before modeling."""
from __future__ import annotations

import csv
from pathlib import Path

from .contracts import validate
from .io import canonical_root, file_hash, read_json, safe_path, now, object_hash


def audit_inputs(root: Path) -> dict:
    root = canonical_root(root)
    manifest_path = safe_path(root, "input_manifest.json")
    manifest = validate("input_manifest", read_json(manifest_path), root=root)
    files, issues = [], []
    for item in manifest["files"]:
        path = safe_path(root, item["path"])
        entry = {"input_id": item["input_id"], "sha256": item["sha256"], "format": path.suffix.lower().lstrip(".") or "unknown",
            "rows": None, "fields": [], "missing_values": None, "duplicate_rows": None, "observations": []}
        try:
            rows = None
            if path.suffix.lower() == ".json":
                value = read_json(path)
                rows = value if isinstance(value, list) else value.get("rows") if isinstance(value, dict) else None
                if rows is None:
                    entry["observations"].append("Structured non-tabular JSON requires task-specific schema interpretation.")
                    if isinstance(value, dict):
                        entry["fields"] = sorted(value)
            elif path.suffix.lower() == ".csv":
                with path.open(encoding="utf-8-sig", newline="") as handle:
                    reader = csv.DictReader(handle)
                    if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
                        raise ValueError("Missing or duplicate CSV column names")
                    rows = list(reader)
            elif path.suffix.lower() in {".txt", ".md", ".tex"}:
                text = path.read_text(encoding="utf-8-sig")
                entry["observations"].append(f"UTF-8 document characters: {len(text)}")
                if not text.strip():
                    raise ValueError("Empty document text")
            elif path.suffix.lower() == ".pdf":
                import fitz
                with fitz.open(path) as document:
                    text = "".join(page.get_text() for page in document)
                    entry["observations"].append(f"PDF pages: {len(document)}; extracted characters: {len(text)}")
                if not text.strip():
                    raise ValueError("PDF has no extractable text; an explicit OCR review is required")
            elif path.suffix.lower() == ".docx":
                from docx import Document
                document = Document(path)
                entry["observations"].append(f"DOCX paragraphs: {len(document.paragraphs)}; tables: {len(document.tables)}")
            else:
                issues.append({"severity": "warning", "input_id": item["input_id"], "message": "No generic structural reader; supply a task-specific input audit"})
            if rows is not None:
                if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
                    raise ValueError("Tabular data requires a list of record objects")
                fields = sorted({str(key) for row in rows for key in row if key is not None})
                if any(None in row for row in rows):
                    raise ValueError("CSV row has undeclared extra columns")
                entry.update(rows=len(rows), fields=fields,
                    missing_values=sum(row.get(field) in (None, "") for row in rows for field in fields),
                    duplicate_rows=len(rows) - len({object_hash(row) for row in rows}))
                for field in fields:
                    values = [row[field] for row in rows if field in row and row[field] is not None]
                    kinds = sorted({type(value).__name__ for value in values})
                    entry["observations"].append(f"{field}: types={','.join(kinds)}")
                    numeric = [value for value in values if type(value) in {int, float}]
                    if numeric:
                        entry["observations"].append(f"{field}: numeric range [{min(numeric)}, {max(numeric)}]")
                if not rows or entry["missing_values"] or entry["duplicate_rows"]:
                    issues.append({"severity": "warning", "input_id": item["input_id"], "message": "Empty/missing/duplicate records require an explicit modeling disposition"})
        except (ValueError, OSError, ImportError) as exc:
            issues.append({"severity": "error", "input_id": item["input_id"], "message": str(exc)})
        files.append(entry)
    result = {"schema_version": "2.0", "audit_id": "input-structure-audit", "actor_id": "data-auditor-tool", "created_at": now(),
        "input_manifest_sha256": file_hash(manifest_path), "files": files, "issues": issues,
        "status": "FAIL" if any(item["severity"] == "error" for item in issues) else "WARN" if issues else "PASS"}
    return validate("data_audit", result, root=root)
