"""Actual DOCX cleanup with explicit Windows replacement-error simulations."""
from pathlib import Path

from docx import Document
import pytest

from test_policy import tool
from mathmode.io import file_hash


@pytest.mark.parametrize("winerror", [5, 32, 33])
def test_scrub_retries_transient_windows_replacement_without_removing_output(tmp_path, monkeypatch, winerror):
    module = tool("build_docx")
    path = tmp_path / "derived.docx"
    document = Document()
    document.core_properties.author = "Synthetic author fixture"
    document.add_paragraph("Synthetic file-replacement regression.")
    document.save(path)
    before = file_hash(path)
    attempts, sleeps = [], []
    replace = Path.replace
    sleep = module.time.sleep
    def pause(seconds):
        sleeps.append(seconds)
        sleep(seconds)
    def temporarily_locked(source, target):
        attempts.append(source)
        if len(attempts) < 3:
            assert file_hash(target) == before, "Denied replacement must preserve the existing package"
            error = PermissionError("Synthetic Windows replacement denial")
            error.winerror = winerror
            raise error
        return replace(source, target)
    monkeypatch.setattr(Path, "replace", temporarily_locked)
    monkeypatch.setattr(module.time, "sleep", pause)
    module.scrub_package(path)
    assert len(attempts) == 3 and sleeps == [0.1, 0.2]
    assert Document(path).core_properties.author == ""
    assert Document(path).paragraphs[0].text == "Synthetic file-replacement regression."
    assert not path.with_suffix(".docx.cleaning").exists()


@pytest.mark.parametrize("winerror,expected_attempts", [(5, 3), (None, 1)])
def test_scrub_preserves_both_packages_and_reports_persistent_denial(tmp_path, monkeypatch, winerror, expected_attempts):
    module = tool("build_docx")
    path = tmp_path / "derived.docx"
    Document().save(path)
    before = file_hash(path)
    attempts = []
    def locked(source, target):
        attempts.append(source)
        error = PermissionError("Synthetic persistent denial")
        if winerror is not None:
            error.winerror = winerror
        raise error
    monkeypatch.setattr(Path, "replace", locked)
    monkeypatch.setattr(module.time, "sleep", lambda seconds: None)
    with pytest.raises(PermissionError, match="persistent denial"):
        module.scrub_package(path)
    assert len(attempts) == expected_attempts
    assert file_hash(path) == before
    Document(path.with_suffix(".docx.cleaning"))  # The completed cleaned package remains readable for diagnosis.
