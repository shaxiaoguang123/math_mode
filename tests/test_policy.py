"""Policy regressions use disposable documentary fixtures, not contest evidence."""
from __future__ import annotations

import copy
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from jsonschema import ValidationError

from mathmode.policy import (
    REPO_ROOT, anonymous_records, audit_policy, file_hash, load_policy,
    page_issues, read_json, toc_settings, validate_policy,
)


def tool(name):
    path = REPO_ROOT / "华为杯_论文规范模板" / "tools" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def paper(tmp_path):
    root = tmp_path / "paper"
    root.mkdir()
    (root / "abstract.tex").write_text("An engineering policy fixture.\n", encoding="utf-8")
    (root / "q1.tex").write_text(r"\section{Analysis}" + "\n" + r"\text{95\%}" + "\n", encoding="utf-8")
    manifest = {"title": "Policy fixture", "keywords": ["a", "b", "c", "d"],
                "abstract_tex_path": "abstract.tex", "chapters": [{
                    "chapter_id": "q1", "title": "Analysis", "role": "problem",
                    "order": 1, "tex_path": "q1.tex"}]}
    path = root / "manifest.json"
    path.write_text(json.dumps(manifest), encoding="utf-8")
    return root, path, manifest


@pytest.fixture
def policy():
    return load_policy()


def test_default_policy_never_claims_verified_official_rules(policy, tmp_path):
    assert audit_policy(policy, tmp_path)["status"] == "BLOCKED"
    assert policy["official"]["page_policy"]["min_body_pages"] is None


@pytest.mark.parametrize("mode,recommendation,expected", [
    ("required", False, True), ("forbidden", True, False),
    ("optional", False, False), ("optional", True, True),
])
def test_toc_rules_override_project_preferences(policy, mode, recommendation, expected):
    policy["official"]["toc_policy"]["mode"] = mode
    policy["project_recommendations"]["include_toc"] = recommendation
    assert toc_settings(policy)[0] is expected


def test_historical_minimum_cannot_become_official_failure(policy):
    issues = page_issues({"start": 2, "end": 9, "pages": 8}, 10, policy, {"required_body_pages": 45})
    assert issues and all(item["severity"] == "warning" for item in issues)
    assert all(item["blocking"] is False for item in issues)
    policy["official"]["page_policy"]["min_body_pages"] = 12
    assert any(i["severity"] == "error" for i in page_issues(
        {"start": 2, "end": 9, "pages": 8}, 10, policy, {}))


def test_official_limits_fail_closed_when_body_unknown(policy):
    policy["official"]["page_policy"].update(min_body_pages=2, max_total_pages=4)
    issues = page_issues({"start": None, "end": None, "pages": 0}, 6, policy, {})
    assert {i["code"] for i in issues} == {"official_body_scope_unknown", "official_total_pages_exceeded"}


def test_identity_cover_scope_excludes_only_declared_cover(policy):
    policy["official"]["cover_policy"] = {"mode": "identity_cover", "pages": 1}
    policy["official"]["anonymity_policy"]["scope"] = "after_cover"
    records = [{"physical_page": 1, "text": "Team identity"}, {"physical_page": 2, "text": "body"}]
    assert anonymous_records(records, policy) == records[1:]
    policy["official"]["anonymity_policy"]["scope"] = "all"
    assert anonymous_records(records, policy) == records


def test_reviewed_snapshot_hashes_and_scope_are_rechecked(policy, tmp_path):
    # Synthetic source bytes test the mechanical provenance check, not authenticity.
    (tmp_path / "rules.txt").write_text("Synthetic policy test document", encoding="utf-8")
    (tmp_path / "template.txt").write_text("Synthetic template", encoding="utf-8")
    policy["year"] = 2026
    policy["verification"].update(status="verified", reviewed_by="engineering-fixture",
                                  reviewed_at="2026-09-08T00:00:00Z")
    policy["official_template_path"] = "template.txt"
    policy["official_template_sha256"] = file_hash(tmp_path / "template.txt")
    policy["official_sources"] = [{"source_id": "test", "url": "https://example.org/rules",
        "authority": "Synthetic fixture", "checked_at": "2026-09-08T00:00:00Z",
        "snapshot_path": "rules.txt", "sha256": file_hash(tmp_path / "rules.txt"),
        "scope": ["competition", "year", "edition", "template", *policy["official"]]}]
    assert audit_policy(policy, tmp_path)["status"] == "PASS"
    policy["official_sources"][0]["scope"].remove("cover_policy")
    assert audit_policy(policy, tmp_path)["status"] == "BLOCKED"
    policy["official_sources"][0]["scope"].append("cover_policy")
    (tmp_path / "rules.txt").write_text("changed", encoding="utf-8")
    assert audit_policy(policy, tmp_path)["status"] == "BLOCKED"


@pytest.mark.parametrize("content", ['{"x":NaN}', '{"x":Infinity}', '{"x":1,"x":2}'])
def test_invalid_json_never_enters_policy(content, tmp_path):
    path = tmp_path / "bad.json"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(ValueError):
        read_json(path)


def test_policy_rejects_template_code_in_edition_and_invalid_cover(policy):
    broken = copy.deepcopy(policy)
    broken["edition"] = r"二十四}\input{secret}"
    with pytest.raises(ValidationError):
        validate_policy(broken)
    policy["official"]["cover_policy"]["pages"] = 1
    with pytest.raises(ValueError):
        validate_policy(policy)


def test_tex_builder_and_auditor_agree_on_toc_and_new_edition(paper, policy):
    root, manifest_path, manifest = paper
    policy["edition"] = "二十四"
    builder, auditor = tool("build_latex"), tool("audit_tex")
    for include in (False, True):
        policy["project_recommendations"]["include_toc"] = include
        policy["official"]["toc_policy"]["depth"] = 2
        text, _ = builder.build_main(manifest, root, policy)
        assert r"\GMCMContestTitle{二十四}" in text
        assert (r"\maketoc" in text) is include
        (root / "paper.tex").write_text(text, encoding="utf-8")
        report = auditor.audit(manifest_path, root / "paper.tex", policy)
        assert report["status"] == "PASS", report["failures"]
        rendered = auditor.audit(manifest_path, root / "paper.tex", policy, stage="render")
        assert (rendered["status"] == "FAIL") is include  # missing .toc must not pass


def test_tex_cover_identity_allowed_but_body_leak_is_blocked(paper, policy):
    root, manifest_path, manifest = paper
    policy["official"]["cover_policy"] = {"mode": "identity_cover", "pages": 1}
    policy["official"]["anonymity_policy"].update(scope="after_cover", identity_terms=["FixturePerson"])
    (root / "cover.tex").write_text("FixturePerson", encoding="utf-8")
    builder, auditor = tool("build_latex"), tool("audit_tex")
    with pytest.raises(ValueError):
        builder.build_main(manifest, root, policy)
    text, inputs = builder.build_main(manifest, root, policy, "cover.tex")
    assert inputs[0] == "cover.tex"
    (root / "paper.tex").write_text(text, encoding="utf-8")
    assert auditor.audit(manifest_path, root / "paper.tex", policy, "cover.tex")["status"] == "PASS"
    (root / "q1.tex").write_text("FixturePerson", encoding="utf-8")
    result = auditor.audit(manifest_path, root / "paper.tex", policy, "cover.tex")
    assert any(f["code"] == "anonymous_source" for f in result["failures"])
    (root / "q1.tex").write_text(r"\input{./cover.tex}", encoding="utf-8")
    result = auditor.audit(manifest_path, root / "paper.tex", policy, "cover.tex")
    assert any(f["code"] == "cover_not_reused_in_body" for f in result["failures"])


def test_tex_comment_and_prefix_brace_regressions():
    auditor = tool("audit_tex")
    assert auditor.brace_balance(auditor.remove_comments(r"\text{95\%} % comment")) == 0
    assert auditor.remove_comments(r"a\\%comment") == r"a\\"
    assert auditor.brace_balance("}{") != 0


def test_pdf_heading_and_page_regex_regressions():
    auditor = tool("audit_paper")
    assert auditor.REFERENCE_RE.fullmatch("参考文献")
    assert auditor.REFERENCE_RE.fullmatch(" References ")
    assert auditor.APPENDIX_RE.fullmatch("Appendix A")
    assert auditor.PAGE_RE.fullmatch("12")
    assert not auditor.PAGE_RE.fullmatch("1234")
    records = [{"physical_page": 2, "text": "Abstract\nContent"},
               {"physical_page": 3, "text": "Continuation\nKeywords: policy"}]
    assert auditor.abstract_page_count(records) == 2
    assert auditor.abstract_page_count(records[:1]) is None


@pytest.mark.parametrize("include", [False, True])
def test_docx_derived_toc_and_edition_match_policy(paper, policy, include, tmp_path):
    root, manifest_path, _ = paper
    policy["edition"] = "二十五"
    policy["project_recommendations"]["include_toc"] = include
    policy["official"]["toc_policy"]["depth"] = 2
    policy_path = tmp_path / "policy.json"
    policy_path.write_text(json.dumps(policy), encoding="utf-8")
    template = tmp_path / "template.docx"
    Document().save(template)
    original = file_hash(template)
    args = SimpleNamespace(input=manifest_path, template=template, output=root / "derived.docx",
                           manifest_out=None, policy=policy_path, cover_tex=None)
    tool("build_docx").build(args)
    assert file_hash(template) == original
    text = "\n".join(p.text for p in Document(args.output).paragraphs)
    assert "第二十五届" in text
    assert ("目录" in text) is include
    result = tool("audit_docx").audit(args.output, policy)
    assert result["status"] == "PASS", result["failures"]


def test_docx_anonymity_scans_table_cells(paper, policy, tmp_path):
    root, manifest_path, _ = paper
    template = tmp_path / "template.docx"
    Document().save(template)
    args = SimpleNamespace(input=manifest_path, template=template, output=root / "derived.docx",
                           manifest_out=None, policy=None, cover_tex=None)
    tool("build_docx").build(args)
    document = Document(args.output)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "FixturePerson"
    document.save(args.output)
    policy["official"]["anonymity_policy"]["identity_terms"] = ["FixturePerson"]
    result = tool("audit_docx").audit(args.output, policy)
    assert any(f["code"] == "identity_text" for f in result["failures"])


def test_docx_identity_cover_keeps_body_scope_and_template_intact(paper, policy, tmp_path):
    root, manifest_path, _ = paper
    template = tmp_path / "template.docx"
    Document().save(template)
    original = file_hash(template)
    policy["official"]["cover_policy"] = {"mode": "identity_cover", "pages": 1}
    policy["official"]["anonymity_policy"].update(scope="after_cover", identity_terms=["FixturePerson"])
    policy_path = tmp_path / "policy.json"
    policy_path.write_text(json.dumps(policy), encoding="utf-8")
    (root / "cover.tex").write_text("FixturePerson", encoding="utf-8")
    args = SimpleNamespace(input=manifest_path, template=template, output=root / "derived.docx",
                           manifest_out=None, policy=policy_path, cover_tex="cover.tex")
    builder, auditor = tool("build_docx"), tool("audit_docx")
    builder.build(args)
    assert file_hash(template) == original
    result = auditor.audit(args.output, policy)
    assert result["status"] == "PASS", result["failures"]
    document = Document(args.output)
    document.add_paragraph("FixturePerson")
    document.save(args.output)
    assert auditor.audit(args.output, policy)["status"] == "FAIL"
    args.output = template
    with pytest.raises(ValueError, match="original DOCX"):
        builder.build(args)
